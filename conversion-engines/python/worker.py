#!/usr/bin/env python3
"""Offline conversion worker. Requests and progress use JSON lines over stdio."""
import json
import re
import sys
import traceback
from pathlib import Path


def emit(event, **values):
    # Escape non-ASCII characters so the JSON-lines protocol stays ASCII even
    # when a frozen Windows worker inherits a legacy console code page.
    print(json.dumps({"event": event, **values}, ensure_ascii=True), flush=True)


def unique_path(directory, stem, suffix):
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    target = directory / f"{stem}{suffix}"
    index = 2
    while target.exists():
        target = directory / f"{stem} ({index}){suffix}"
        index += 1
    return target


def destination(source, request, suffix):
    source = Path(source)
    return unique_path(request.get("outputDir") or source.parent, output_stem(request, source, suffix), suffix)


def output_stem(request, source, suffix=""):
    source = Path(source)
    options = request.get("options", {})
    custom_name = str(options.get("outputNamesByPath", {}).get(str(source)) or options.get("outputName") or "").strip()
    if not custom_name:
        return source.stem
    stem = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", custom_name).strip(" .")
    if suffix and stem.lower().endswith(suffix.lower()):
        stem = stem[:-len(suffix)].rstrip(" .")
    return stem or source.stem


def encode_image_to_target(image, pil_format, target_bytes, quality):
    from io import BytesIO
    from PIL import Image

    def encode(candidate, candidate_quality=quality):
        buffer = BytesIO()
        options = {"optimize": True}
        if pil_format in ("JPEG", "WEBP"):
            options["quality"] = candidate_quality
        if pil_format == "WEBP":
            options["method"] = 6
        if pil_format == "PNG":
            options["compress_level"] = 9
        candidate.save(buffer, format=pil_format, **options)
        return buffer.getvalue()

    if pil_format in ("JPEG", "WEBP"):
        low, high, best = 1, quality, None
        while low <= high:
            candidate_quality = (low + high) // 2
            data = encode(image, candidate_quality)
            if len(data) <= target_bytes:
                best = data
                low = candidate_quality + 1
            else:
                high = candidate_quality - 1
        if best is None:
            raise ValueError("目标体积过小，当前图片即使使用最低质量也无法达到")
        return best

    if pil_format == "PNG":
        low, high, best = 0.03, 1.0, None
        original_size = image.size
        for _ in range(16):
            scale = (low + high) / 2
            size = (max(1, round(original_size[0] * scale)), max(1, round(original_size[1] * scale)))
            candidate = image if size == original_size else image.resize(size, Image.Resampling.LANCZOS)
            data = encode(candidate)
            if len(data) <= target_bytes:
                best = data
                low = scale
            else:
                high = scale
        if best is None:
            raise ValueError("目标体积过小，PNG 图片缩小后仍无法达到")
        return best

    raise ValueError(f"{pil_format} 暂不支持按目标体积压缩")


def image_convert(request):
    from PIL import Image
    source = Path(request["path"])
    options = request.get("options", {})
    target = str(options.get("format", "PNG")).upper()
    formats = {"JPG": ("JPEG", ".jpg"), "JPEG": ("JPEG", ".jpg"), "PNG": ("PNG", ".png"), "WEBP": ("WEBP", ".webp"), "ICO": ("ICO", ".ico")}
    if target not in formats:
        raise ValueError(f"不支持的图片格式：{target}")
    pil_format, suffix = formats[target]
    output = destination(source, request, suffix)
    emit("progress", percent=5, message="读取图片")
    with Image.open(source) as original:
        image = original.copy()
        if pil_format == "ICO":
            image = image.convert("RGBA")
            image.save(output, format="ICO", sizes=[(16, 16), (32, 32), (48, 48), (64, 64), (128, 128), (256, 256)])
        else:
            if pil_format == "JPEG" and image.mode not in ("RGB", "L"):
                rgba = image.convert("RGBA")
                background = Image.new("RGB", image.size, (255, 255, 255))
                background.paste(rgba, mask=rgba.getchannel("A"))
                image = background
            if request["operation"] == "image.resize":
                resize = options.get("resize", {})
                if resize.get("mode") == "width":
                    width = max(1, int(resize.get("width", 800)))
                    height = max(1, round(image.height * width / image.width))
                elif resize.get("mode") == "height":
                    height = max(1, int(resize.get("height", 600)))
                    width = max(1, round(image.width * height / image.height))
                else:
                    percent = max(1, int(resize.get("percent", 50)))
                    width, height = max(1, round(image.width * percent / 100)), max(1, round(image.height * percent / 100))
                image = image.resize((width, height), Image.Resampling.LANCZOS)
            emit("progress", percent=70, message="编码图片")
            quality = max(1, min(100, int(options.get("quality", 85))))
            save_options = {"quality": quality}
            if pil_format in ("JPEG", "PNG"):
                save_options["optimize"] = True
            elif pil_format == "WEBP":
                save_options["method"] = 6
            target_map = options.get("targetBytesByPath", {})
            target_bytes = target_map.get(str(source)) if isinstance(target_map, dict) else None
            if request["operation"] == "image.compress" and target_bytes is not None:
                encoded = encode_image_to_target(image, pil_format, max(1, int(target_bytes)), quality)
                output.write_bytes(encoded)
            else:
                image.save(output, format=pil_format, **save_options)
    emit("progress", percent=95, message="写入输出文件")
    return {"outputs": [str(output)]}


def pdf_to_images(request):
    from pdf2image import convert_from_path
    from PyPDF2 import PdfReader
    source = Path(request["path"])
    options = request.get("options", {})
    dpi = int(options.get("dpi", 160))
    fmt = options.get("format", "PNG").lower()
    if fmt == "jpg":
        fmt = "jpeg"
    poppler = request.get("popplerPath")
    reader = PdfReader(str(source), strict=False)
    if reader.is_encrypted and not reader.decrypt(""):
        raise ValueError("PDF 已加密，当前不支持无密码读取")
    page_count = len(reader.pages)
    if page_count < 1:
        raise ValueError("PDF 没有可导出的页面")
    folder = Path(request.get("outputDir") or source.parent)
    outputs = []
    suffix = ".jpg" if options.get("format", "PNG").upper() in ("JPG", "JPEG") else ".png"
    for index in range(1, page_count + 1):
        image = convert_from_path(str(source), dpi=dpi, fmt=fmt, poppler_path=poppler, first_page=index, last_page=index)[0]
        out = unique_path(folder, f"{output_stem(request, source, suffix)}-{index:03d}", suffix)
        if suffix == ".jpg" and image.mode != "RGB":
            image = image.convert("RGB")
        image.save(out)
        outputs.append(str(out))
        emit("progress", percent=round(index * 100 / max(1, page_count)), message=f"导出第 {index}/{page_count} 页")
    return {"outputs": outputs}


def images_to_pdf(request):
    from PIL import Image
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    files = request.get("paths", [])
    if not files:
        raise ValueError("未选择图片")
    first = Path(files[0])
    folder = Path(request.get("outputDir") or first.parent)
    custom_name = request.get("options", {}).get("outputNamesByPath", {}).get(str(first)) or request.get("options", {}).get("outputName")
    out = unique_path(folder, output_stem(request, first, ".pdf") if custom_name else "merged-images", ".pdf")
    pdf = canvas.Canvas(str(out), pageCompression=1)
    for index, file_path in enumerate(files, 1):
        with Image.open(file_path) as image:
            width, height = image.size
            pdf.setPageSize((width, height))
            rgba = image.convert("RGBA")
            white_background = Image.new("RGB", rgba.size, (255, 255, 255))
            white_background.paste(rgba, mask=rgba.getchannel("A"))
            pdf.drawImage(ImageReader(white_background), 0, 0, width=width, height=height)
            pdf.showPage()
        emit("progress", percent=round(index * 100 / len(files)), message=f"合并第 {index}/{len(files)} 张图片")
    pdf.save()
    return {"outputs": [str(out)]}


def excel_convert(request):
    import pandas as pd
    source = Path(request["path"])
    operation = request["operation"]
    folder = Path(request.get("outputDir") or source.parent)
    folder.mkdir(parents=True, exist_ok=True)
    emit("progress", percent=10, message="读取表格")
    if operation == "excel.xlsx-to-csv":
        frame = pd.read_excel(source, engine="openpyxl")
        out = unique_path(folder, output_stem(request, source, ".csv"), ".csv")
        frame.to_csv(out, index=False, encoding="utf-8-sig")
    elif operation == "excel.csv-to-xlsx":
        frame = pd.read_csv(source)
        out = unique_path(folder, output_stem(request, source, ".xlsx"), ".xlsx")
        frame.to_excel(out, index=False, engine="openpyxl")
    elif operation == "excel.xlsx-to-json":
        frame = pd.read_excel(source, engine="openpyxl")
        out = unique_path(folder, output_stem(request, source, ".json"), ".json")
        out.write_text(frame.to_json(orient="records", force_ascii=False, indent=2), encoding="utf-8")
    elif operation == "excel.json-to-xlsx":
        values = json.loads(source.read_text(encoding="utf-8"))
        if isinstance(values, dict):
            values = values if all(isinstance(value, list) for value in values.values()) else [values]
        if not isinstance(values, list):
            raise ValueError("JSON 根节点必须是数组或对象")
        out = unique_path(folder, output_stem(request, source, ".xlsx"), ".xlsx")
        pd.DataFrame(values).to_excel(out, index=False, engine="openpyxl")
    else:
        raise ValueError(f"不支持的 Excel 操作：{operation}")
    emit("progress", percent=90, message="保存转换结果")
    return {"outputs": [str(out)]}


def docx_convert(request):
    from docx import Document
    source = Path(request["path"])
    folder = Path(request.get("outputDir") or source.parent)
    folder.mkdir(parents=True, exist_ok=True)
    emit("progress", percent=10, message="读取 Word 文档")
    if request["operation"] == "word.docx-to-txt":
        document = Document(str(source))
        lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                lines.append("\t".join(cell.text.replace("\n", " ") for cell in row.cells))
        out = unique_path(folder, output_stem(request, source, ".txt"), ".txt")
        out.write_text("\n".join(lines), encoding="utf-8")
    else:
        import mammoth
        from markdownify import markdownify
        with source.open("rb") as docx:
            html = mammoth.convert_to_html(docx).value
        out = unique_path(folder, output_stem(request, source, ".md"), ".md")
        out.write_text(markdownify(html, heading_style="ATX", bullets="-").rstrip() + "\n", encoding="utf-8")
    emit("progress", percent=90, message="保存转换结果")
    return {"outputs": [str(out)]}


def run(request):
    operation = request["operation"]
    if operation.startswith("image."):
        return image_convert(request)
    if operation == "pdf.pdf-to-images":
        return pdf_to_images(request)
    if operation == "pdf.images-to-pdf":
        return images_to_pdf(request)
    if operation.startswith("excel."):
        return excel_convert(request)
    if operation.startswith("word."):
        return docx_convert(request)
    raise ValueError(f"未知转换操作：{operation}")


if __name__ == "__main__":
    try:
        # Electron writes the JSON-lines protocol as UTF-8. Decode the bytes
        # explicitly instead of relying on Windows' active code page, so
        # Chinese and other Unicode file paths survive intact.
        request_line = sys.stdin.buffer.readline().decode("utf-8")
        request = json.loads(request_line)
        emit("result", **run(request))
    except Exception as error:
        emit("error", message=str(error), detail=traceback.format_exc())
        sys.exit(1)
